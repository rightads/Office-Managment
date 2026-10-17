from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.schemas.attendance_reports import (
    AttendanceTodayReportResponse,
    EmployeeMonthlyReportResponse,
    EmployeeYearlyReportResponse,
    MonthlyOverviewResponse,
    ReportType,
)


class AttendanceExportService:
    """Service generating downloadable Excel (.xlsx) and PDF (.pdf) reports."""

    # -------------------------------------------------------------------------
    # EXCEL EXPORT (OpenPyXL)
    # -------------------------------------------------------------------------
    @classmethod
    def generate_excel(cls, report_type: ReportType, data: Any) -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None

        # Color schemes & fonts
        navy_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        accent_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        white_bold_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        section_font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        bold_font = Font(name="Calibri", size=10, bold=True)
        regular_font = Font(name="Calibri", size=10)

        thin_side = Side(border_style="thin", color="CBD5E1")
        grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        if report_type == "today":
            assert isinstance(data, AttendanceTodayReportResponse)
            ws.title = "Today Attendance"
            ws["A1"] = "RIGHT ADS - TODAY ATTENDANCE REPORT"
            ws["A1"].font = title_font
            ws["A2"] = (
                f"Date: {data.date.isoformat()} | Total Active Employees: {data.total_active_employees}"
            )
            ws["A2"].font = bold_font

            # KPI Summary Table
            kpis = [
                ("Present", data.present_count),
                ("Absent", data.absent_count),
                ("Late", data.late_count),
                ("Half Day", data.half_day_count),
                ("On Leave", data.on_leave_count),
            ]
            row_idx = 4
            ws.cell(row=row_idx, column=1, value="Metric").font = white_bold_font
            ws.cell(row=row_idx, column=1).fill = navy_header_fill
            ws.cell(row=row_idx, column=2, value="Count").font = white_bold_font
            ws.cell(row=row_idx, column=2).fill = navy_header_fill
            for kpi_label, kpi_val in kpis:
                row_idx += 1
                c1 = ws.cell(row=row_idx, column=1, value=kpi_label)
                c2 = ws.cell(row=row_idx, column=2, value=kpi_val)
                c1.font = regular_font
                c2.font = bold_font
                c1.border = grid_border
                c2.border = grid_border

            # Employee Detail Table
            row_idx += 2
            ws.cell(row=row_idx, column=1, value="EMPLOYEE ROSTER BREAKDOWN").font = section_font
            row_idx += 1
            headers = ["Employee Code", "Full Name", "Department", "Designation", "Status"]
            for col_num, h_text in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_num, value=h_text)
                cell.font = white_bold_font
                cell.fill = navy_header_fill
                cell.alignment = center_align
                cell.border = grid_border

            for status_group in data.status_breakdown:
                for emp in status_group.employees:
                    row_idx += 1
                    ws.cell(row=row_idx, column=1, value=emp.employee_code).alignment = left_align
                    ws.cell(row=row_idx, column=2, value=emp.full_name).alignment = left_align
                    ws.cell(
                        row=row_idx, column=3, value=emp.department_name or "N/A"
                    ).alignment = left_align
                    ws.cell(
                        row=row_idx, column=4, value=emp.designation or "N/A"
                    ).alignment = left_align
                    st_cell = ws.cell(row=row_idx, column=5, value=status_group.status_name)
                    st_cell.alignment = center_align
                    for col_num in range(1, 6):
                        c = ws.cell(row=row_idx, column=col_num)
                        c.font = regular_font
                        c.border = grid_border

        elif report_type == "employee_monthly":
            assert isinstance(data, EmployeeMonthlyReportResponse)
            ws.title = f"{data.month:02d}-{data.year} Attendance"
            ws["A1"] = f"RIGHT ADS - MONTHLY ATTENDANCE REPORT ({data.month:02d}/{data.year})"
            ws["A1"].font = title_font
            ws["A2"] = (
                f"Employee: {data.employee_name} ({data.employee_code}) | Dept: {data.department_name or 'N/A'}"
            )
            ws["A2"].font = bold_font

            # KPI Block
            metrics = [
                ("Working Days", data.working_days),
                ("Present Days", data.present_days),
                ("Absent Days", data.absent_days),
                ("Half Days", data.half_days),
                ("Late Days", data.late_days),
                ("Leave Days", data.leave_days),
                ("Total Hours Worked", f"{data.total_hours_worked:.2f} hrs"),
                ("Average Hours / Day", f"{data.average_hours_per_day:.2f} hrs"),
                ("Attendance %", f"{data.attendance_percentage:.1f}%"),
            ]
            row_idx = 4
            for m_label, m_val in metrics:
                c1 = ws.cell(row=row_idx, column=1, value=m_label)
                c2 = ws.cell(row=row_idx, column=2, value=m_val)
                c1.font = bold_font
                c2.font = regular_font
                c1.fill = accent_fill
                c1.border = grid_border
                c2.border = grid_border
                row_idx += 1

            row_idx += 1
            ws.cell(row=row_idx, column=1, value="DAILY LOG BREAKDOWN").font = section_font
            row_idx += 1

            headers = [
                "Date",
                "Day",
                "Status",
                "Check-In",
                "Check-Out",
                "Hours Worked",
                "Overtime",
                "Late?",
                "Half Day?",
                "Manual?",
            ]
            for col_num, h_text in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_num, value=h_text)
                cell.font = white_bold_font
                cell.fill = navy_header_fill
                cell.alignment = center_align
                cell.border = grid_border

            for day_rec in data.daily_records:
                row_idx += 1
                ws.cell(
                    row=row_idx, column=1, value=day_rec.date.isoformat()
                ).alignment = center_align
                ws.cell(row=row_idx, column=2, value=day_rec.day_of_week).alignment = left_align
                ws.cell(
                    row=row_idx, column=3, value=day_rec.status_name or "N/A"
                ).alignment = center_align
                ws.cell(
                    row=row_idx,
                    column=4,
                    value=day_rec.first_check_in.strftime("%H:%M:%S")
                    if day_rec.first_check_in
                    else "-",
                ).alignment = center_align
                ws.cell(
                    row=row_idx,
                    column=5,
                    value=day_rec.last_check_out.strftime("%H:%M:%S")
                    if day_rec.last_check_out
                    else "-",
                ).alignment = center_align
                ws.cell(
                    row=row_idx, column=6, value=day_rec.total_hours_worked
                ).alignment = right_align
                ws.cell(row=row_idx, column=7, value=day_rec.overtime_hours).alignment = right_align
                ws.cell(
                    row=row_idx, column=8, value="Yes" if day_rec.is_late else "No"
                ).alignment = center_align
                ws.cell(
                    row=row_idx, column=9, value="Yes" if day_rec.is_half_day else "No"
                ).alignment = center_align
                ws.cell(
                    row=row_idx, column=10, value="Yes" if day_rec.correction_applied else "No"
                ).alignment = center_align

                for col_num in range(1, 11):
                    c = ws.cell(row=row_idx, column=col_num)
                    c.font = regular_font
                    c.border = grid_border

        elif report_type == "employee_yearly":
            assert isinstance(data, EmployeeYearlyReportResponse)
            ws.title = f"Year {data.year} Attendance"
            ws["A1"] = f"RIGHT ADS - YEARLY ATTENDANCE REPORT ({data.year})"
            ws["A1"].font = title_font
            ws["A2"] = (
                f"Employee: {data.employee_name} ({data.employee_code}) | Dept: {data.department_name or 'N/A'}"
            )
            ws["A2"].font = bold_font

            # KPI Block
            metrics = [
                ("Annual Working Days", data.annual_working_days),
                ("Annual Present Days", data.annual_present_days),
                ("Annual Absent Days", data.annual_absent_days),
                ("Annual Leave Days", data.annual_leave_days),
                ("Annual Total Hours", f"{data.annual_total_hours:.2f} hrs"),
                ("Annual Attendance %", f"{data.annual_attendance_percentage:.1f}%"),
            ]
            row_idx = 4
            for ym_label, ym_val in metrics:
                c1 = ws.cell(row=row_idx, column=1, value=ym_label)
                c2 = ws.cell(row=row_idx, column=2, value=ym_val)
                c1.font = bold_font
                c2.font = regular_font
                c1.fill = accent_fill
                c1.border = grid_border
                c2.border = grid_border
                row_idx += 1

            row_idx += 1
            ws.cell(row=row_idx, column=1, value="12-MONTH SUMMARY").font = section_font
            row_idx += 1

            headers = [
                "Month",
                "Working Days",
                "Present",
                "Absent",
                "Half Days",
                "Late",
                "Leave",
                "Holidays",
                "Total Hours",
                "Attendance %",
            ]
            for col_num, h_text in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_num, value=h_text)
                cell.font = white_bold_font
                cell.fill = navy_header_fill
                cell.alignment = center_align
                cell.border = grid_border

            for m in data.months:
                row_idx += 1
                ws.cell(row=row_idx, column=1, value=m.month_name).alignment = left_align
                ws.cell(row=row_idx, column=2, value=m.working_days).alignment = right_align
                ws.cell(row=row_idx, column=3, value=m.present_days).alignment = right_align
                ws.cell(row=row_idx, column=4, value=m.absent_days).alignment = right_align
                ws.cell(row=row_idx, column=5, value=m.half_days).alignment = right_align
                ws.cell(row=row_idx, column=6, value=m.late_days).alignment = right_align
                ws.cell(row=row_idx, column=7, value=m.leave_days).alignment = right_align
                ws.cell(row=row_idx, column=8, value=m.holiday_days).alignment = right_align
                ws.cell(row=row_idx, column=9, value=m.total_hours_worked).alignment = right_align
                ws.cell(
                    row=row_idx, column=10, value=f"{m.attendance_percentage:.1f}%"
                ).alignment = right_align

                for col_num in range(1, 11):
                    c = ws.cell(row=row_idx, column=col_num)
                    c.font = regular_font
                    c.border = grid_border

        elif report_type == "monthly_overview":
            assert isinstance(data, MonthlyOverviewResponse)
            ws.title = f"Overview {data.month:02d}-{data.year}"
            ws["A1"] = f"RIGHT ADS - MONTHLY ATTENDANCE OVERVIEW ({data.month:02d}/{data.year})"
            ws["A1"].font = title_font
            ws["A2"] = f"Total Employees: {data.total_employees}"
            ws["A2"].font = bold_font

            row_idx = 4
            days_in_month = 31
            if data.rows:
                days_in_month = max(max(r.days.keys()) if r.days else 30 for r in data.rows)

            headers = ["Code", "Employee Name", "Department"]
            for d in range(1, days_in_month + 1):
                headers.append(str(d))
            headers.extend(["Present", "Absent", "Half Day", "Late", "Leave", "Total Hrs", "Att %"])

            for col_num, h_text in enumerate(headers, 1):
                cell = ws.cell(row=row_idx, column=col_num, value=h_text)
                cell.font = white_bold_font
                cell.fill = navy_header_fill
                cell.alignment = center_align
                cell.border = grid_border

            for row_data in data.rows:
                row_idx += 1
                col_idx = 1
                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.employee_code
                ).alignment = left_align
                col_idx += 1
                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.employee_name
                ).alignment = left_align
                col_idx += 1
                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.department_name or "N/A"
                ).alignment = left_align
                col_idx += 1

                for d in range(1, days_in_month + 1):
                    day_status = row_data.days.get(d, "-")
                    # Shorten code for clean grid display
                    short_val = {
                        "PRESENT": "P",
                        "ABSENT": "A",
                        "HALF_DAY": "HD",
                        "LATE": "L",
                        "ON_LEAVE": "OL",
                        "HOLIDAY": "H",
                        "WEEK_OFF": "WO",
                        "UPCOMING": "-",
                    }.get(day_status, day_status[:2])
                    c = ws.cell(row=row_idx, column=col_idx, value=short_val)
                    c.alignment = center_align
                    col_idx += 1

                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.present_count
                ).alignment = right_align
                col_idx += 1
                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.absent_count
                ).alignment = right_align
                col_idx += 1
                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.half_day_count
                ).alignment = right_align
                col_idx += 1
                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.late_count
                ).alignment = right_align
                col_idx += 1
                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.leave_count
                ).alignment = right_align
                col_idx += 1
                ws.cell(
                    row=row_idx, column=col_idx, value=row_data.total_hours
                ).alignment = right_align
                col_idx += 1
                ws.cell(
                    row=row_idx, column=col_idx, value=f"{row_data.attendance_percentage:.1f}%"
                ).alignment = right_align

                for c_num in range(1, col_idx + 1):
                    cell = ws.cell(row=row_idx, column=c_num)
                    cell.font = regular_font
                    cell.border = grid_border

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            first_cell = col[0]
            col_letter = get_column_letter(first_cell.column)
            for cell in col:
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 35)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    # -------------------------------------------------------------------------
    # PDF EXPORT (ReportLab)
    # -------------------------------------------------------------------------
    @classmethod
    def generate_pdf(cls, report_type: ReportType, data: Any) -> io.BytesIO:
        stream = io.BytesIO()

        # Landscape for monthly overview matrix, Portrait for all others
        page_size = landscape(A4) if report_type == "monthly_overview" else A4
        margin = 30 if report_type == "monthly_overview" else 36

        doc = SimpleDocTemplate(
            stream,
            pagesize=page_size,
            leftMargin=margin,
            rightMargin=margin,
            topMargin=margin,
            bottomMargin=margin,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#1E3A8A"),
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#475569"),
            spaceAfter=12,
        )
        table_header_style = ParagraphStyle(
            "TableHeader",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.white,
            alignment=1,  # Center
            fontName="Helvetica-Bold",
        )
        table_cell_style = ParagraphStyle(
            "TableCell",
            parent=styles["Normal"],
            fontSize=7,
            leading=9,
            textColor=colors.HexColor("#1E293B"),
            alignment=1,  # Center
        )
        kpi_title_style = ParagraphStyle(
            "KPITitle",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#64748B"),
            alignment=1,
            fontName="Helvetica-Bold",
        )
        kpi_val_style = ParagraphStyle(
            "KPIVal",
            parent=styles["Normal"],
            fontSize=12,
            leading=14,
            textColor=colors.HexColor("#0F172A"),
            alignment=1,
            fontName="Helvetica-Bold",
        )

        story: list[Any] = []

        if report_type == "today":
            assert isinstance(data, AttendanceTodayReportResponse)
            story.append(Paragraph("RIGHT ADS - TODAY ATTENDANCE REPORT", title_style))
            story.append(
                Paragraph(
                    f"Date: {data.date.isoformat()} | Total Active Employees: {data.total_active_employees}",
                    subtitle_style,
                )
            )

            # KPI Grid
            kpi_data = [
                [
                    Paragraph("PRESENT", kpi_title_style),
                    Paragraph("ABSENT", kpi_title_style),
                    Paragraph("LATE", kpi_title_style),
                    Paragraph("HALF DAY", kpi_title_style),
                    Paragraph("ON LEAVE", kpi_title_style),
                ],
                [
                    Paragraph(str(data.present_count), kpi_val_style),
                    Paragraph(str(data.absent_count), kpi_val_style),
                    Paragraph(str(data.late_count), kpi_val_style),
                    Paragraph(str(data.half_day_count), kpi_val_style),
                    Paragraph(str(data.on_leave_count), kpi_val_style),
                ],
            ]
            kpi_table = Table(kpi_data, colWidths=[100, 100, 100, 100, 100])
            kpi_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.append(kpi_table)
            story.append(Spacer(1, 15))

            # Table of employees
            table_rows = [
                [
                    Paragraph("Code", table_header_style),
                    Paragraph("Full Name", table_header_style),
                    Paragraph("Department", table_header_style),
                    Paragraph("Designation", table_header_style),
                    Paragraph("Status", table_header_style),
                ]
            ]
            for status_group in data.status_breakdown:
                for emp in status_group.employees:
                    table_rows.append(
                        [
                            Paragraph(emp.employee_code, table_cell_style),
                            Paragraph(emp.full_name, table_cell_style),
                            Paragraph(emp.department_name or "N/A", table_cell_style),
                            Paragraph(emp.designation or "N/A", table_cell_style),
                            Paragraph(status_group.status_name, table_cell_style),
                        ]
                    )

            if len(table_rows) > 1:
                t = Table(table_rows, colWidths=[70, 140, 120, 110, 80], repeatRows=1)
                t.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                            (
                                "ROWBACKGROUNDS",
                                (0, 1),
                                (-1, -1),
                                [colors.white, colors.HexColor("#F8FAFC")],
                            ),
                            ("TOPPADDING", (0, 0), (-1, -1), 4),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        ]
                    )
                )
                story.append(t)

        elif report_type == "employee_monthly":
            assert isinstance(data, EmployeeMonthlyReportResponse)
            story.append(
                Paragraph(
                    f"RIGHT ADS - MONTHLY ATTENDANCE REPORT ({data.month:02d}/{data.year})",
                    title_style,
                )
            )
            story.append(
                Paragraph(
                    f"Employee: {data.employee_name} ({data.employee_code}) | Dept: {data.department_name or 'N/A'} | Designation: {data.designation or 'N/A'}",
                    subtitle_style,
                )
            )

            # KPI Grid
            kpi_data = [
                [
                    Paragraph("WORKING DAYS", kpi_title_style),
                    Paragraph("PRESENT", kpi_title_style),
                    Paragraph("ABSENT", kpi_title_style),
                    Paragraph("TOTAL HOURS", kpi_title_style),
                    Paragraph("ATTENDANCE %", kpi_title_style),
                ],
                [
                    Paragraph(str(data.working_days), kpi_val_style),
                    Paragraph(str(data.present_days), kpi_val_style),
                    Paragraph(str(data.absent_days), kpi_val_style),
                    Paragraph(f"{data.total_hours_worked:.1f}h", kpi_val_style),
                    Paragraph(f"{data.attendance_percentage:.1f}%", kpi_val_style),
                ],
            ]
            kpi_table = Table(kpi_data, colWidths=[100, 100, 100, 100, 100])
            kpi_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.append(kpi_table)
            story.append(Spacer(1, 15))

            # Daily Table
            table_rows = [
                [
                    Paragraph("Date", table_header_style),
                    Paragraph("Day", table_header_style),
                    Paragraph("Status", table_header_style),
                    Paragraph("Check-In", table_header_style),
                    Paragraph("Check-Out", table_header_style),
                    Paragraph("Hours", table_header_style),
                    Paragraph("Overtime", table_header_style),
                    Paragraph("Late?", table_header_style),
                ]
            ]
            for d in data.daily_records:
                table_rows.append(
                    [
                        Paragraph(d.date.isoformat(), table_cell_style),
                        Paragraph(d.day_of_week[:3], table_cell_style),
                        Paragraph(d.status_name or "-", table_cell_style),
                        Paragraph(
                            d.first_check_in.strftime("%H:%M") if d.first_check_in else "-",
                            table_cell_style,
                        ),
                        Paragraph(
                            d.last_check_out.strftime("%H:%M") if d.last_check_out else "-",
                            table_cell_style,
                        ),
                        Paragraph(f"{d.total_hours_worked:.1f}", table_cell_style),
                        Paragraph(f"{d.overtime_hours:.1f}", table_cell_style),
                        Paragraph("Yes" if d.is_late else "No", table_cell_style),
                    ]
                )

            t = Table(table_rows, colWidths=[65, 45, 95, 65, 65, 55, 55, 50], repeatRows=1)
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F8FAFC")],
                        ),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.append(t)

        elif report_type == "employee_yearly":
            assert isinstance(data, EmployeeYearlyReportResponse)
            story.append(
                Paragraph(
                    f"RIGHT ADS - YEARLY ATTENDANCE REPORT ({data.year})",
                    title_style,
                )
            )
            story.append(
                Paragraph(
                    f"Employee: {data.employee_name} ({data.employee_code}) | Dept: {data.department_name or 'N/A'}",
                    subtitle_style,
                )
            )

            # KPI Grid
            kpi_data = [
                [
                    Paragraph("WORKING DAYS", kpi_title_style),
                    Paragraph("PRESENT", kpi_title_style),
                    Paragraph("ABSENT", kpi_title_style),
                    Paragraph("LEAVE", kpi_title_style),
                    Paragraph("TOTAL HOURS", kpi_title_style),
                    Paragraph("ATTENDANCE %", kpi_title_style),
                ],
                [
                    Paragraph(str(data.annual_working_days), kpi_val_style),
                    Paragraph(str(data.annual_present_days), kpi_val_style),
                    Paragraph(str(data.annual_absent_days), kpi_val_style),
                    Paragraph(str(data.annual_leave_days), kpi_val_style),
                    Paragraph(f"{data.annual_total_hours:.1f}h", kpi_val_style),
                    Paragraph(f"{data.annual_attendance_percentage:.1f}%", kpi_val_style),
                ],
            ]
            kpi_table = Table(kpi_data, colWidths=[80, 80, 80, 80, 90, 90])
            kpi_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
                        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                        ("TOPPADDING", (0, 0), (-1, -1), 6),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ]
                )
            )
            story.append(kpi_table)
            story.append(Spacer(1, 15))

            # 12-Month Table
            table_rows = [
                [
                    Paragraph("Month", table_header_style),
                    Paragraph("Work Days", table_header_style),
                    Paragraph("Present", table_header_style),
                    Paragraph("Absent", table_header_style),
                    Paragraph("Half Days", table_header_style),
                    Paragraph("Late", table_header_style),
                    Paragraph("Leave", table_header_style),
                    Paragraph("Hours", table_header_style),
                    Paragraph("Att %", table_header_style),
                ]
            ]
            for m in data.months:
                table_rows.append(
                    [
                        Paragraph(m.month_name, table_cell_style),
                        Paragraph(str(m.working_days), table_cell_style),
                        Paragraph(str(m.present_days), table_cell_style),
                        Paragraph(str(m.absent_days), table_cell_style),
                        Paragraph(str(m.half_days), table_cell_style),
                        Paragraph(str(m.late_days), table_cell_style),
                        Paragraph(str(m.leave_days), table_cell_style),
                        Paragraph(f"{m.total_hours_worked:.1f}", table_cell_style),
                        Paragraph(f"{m.attendance_percentage:.1f}%", table_cell_style),
                    ]
                )

            t = Table(table_rows, colWidths=[80, 50, 50, 50, 50, 50, 50, 60, 60], repeatRows=1)
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F8FAFC")],
                        ),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(t)

        elif report_type == "monthly_overview":
            assert isinstance(data, MonthlyOverviewResponse)
            story.append(
                Paragraph(
                    f"RIGHT ADS - MONTHLY ATTENDANCE OVERVIEW ({data.month:02d}/{data.year})",
                    title_style,
                )
            )
            story.append(
                Paragraph(
                    f"Total Employees: {data.total_employees}",
                    subtitle_style,
                )
            )

            # Compact matrix for landscape A4
            # Printable width in landscape A4 is 842 - 60 = ~782pt
            table_rows = [
                [
                    Paragraph("Code", table_header_style),
                    Paragraph("Employee Name", table_header_style),
                    Paragraph("Department", table_header_style),
                    Paragraph("P", table_header_style),
                    Paragraph("A", table_header_style),
                    Paragraph("HD", table_header_style),
                    Paragraph("L", table_header_style),
                    Paragraph("OL", table_header_style),
                    Paragraph("Hours", table_header_style),
                    Paragraph("Att %", table_header_style),
                ]
            ]
            for row in data.rows:
                table_rows.append(
                    [
                        Paragraph(row.employee_code, table_cell_style),
                        Paragraph(row.employee_name, table_cell_style),
                        Paragraph(row.department_name or "-", table_cell_style),
                        Paragraph(str(row.present_count), table_cell_style),
                        Paragraph(str(row.absent_count), table_cell_style),
                        Paragraph(str(row.half_day_count), table_cell_style),
                        Paragraph(str(row.late_count), table_cell_style),
                        Paragraph(str(row.leave_count), table_cell_style),
                        Paragraph(f"{row.total_hours:.1f}", table_cell_style),
                        Paragraph(f"{row.attendance_percentage:.1f}%", table_cell_style),
                    ]
                )

            t = Table(
                table_rows, colWidths=[70, 180, 140, 45, 45, 45, 45, 45, 65, 65], repeatRows=1
            )
            t.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [colors.white, colors.HexColor("#F8FAFC")],
                        ),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(t)

        doc.build(story)
        stream.seek(0)
        return stream
