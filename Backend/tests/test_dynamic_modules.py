from __future__ import annotations

import io
from typing import Any

import openpyxl
import pytest
from httpx import AsyncClient


@pytest.fixture
async def leads_module(
    async_client: AsyncClient,
    admin_auth_headers: dict[str, str],
) -> dict[str, Any]:
    """Helper fixture to create leads dynamic module definition."""
    create_payload = {
        "code": "leads",
        "label": "Leads & Prospects",
        "record_label_singular": "Lead",
        "record_label_plural": "Leads",
        "description": "Sales lead pipeline",
        "fields": [
            {
                "key": "title",
                "label": "Lead Title",
                "type": "text",
                "is_required": True,
                "is_list_visible": True,
                "is_searchable": True,
            },
            {
                "key": "contact_name",
                "label": "Contact Name",
                "type": "text",
                "is_required": True,
                "is_list_visible": True,
                "is_searchable": True,
            },
            {
                "key": "deal_value",
                "label": "Deal Value",
                "type": "currency",
                "is_required": False,
                "is_list_visible": True,
                "is_filterable": True,
            },
            {
                "key": "source",
                "label": "Lead Source",
                "type": "single_select",
                "is_required": False,
                "options": [
                    {"value": "web", "label": "Website", "color": "#3B82F6"},
                    {"value": "referral", "label": "Referral", "color": "#10B981"},
                ],
            },
        ],
        "statuses": [
            {"code": "new", "label": "New", "color": "#3B82F6", "sort_order": 0},
            {"code": "contacted", "label": "Contacted", "color": "#8B5CF6", "sort_order": 1},
            {"code": "won", "label": "Won", "color": "#10B981", "sort_order": 2, "is_terminal": True},
        ],
    }
    res = await async_client.post("/api/v1/modules/definitions", headers=admin_auth_headers, json=create_payload)
    assert res.status_code == 201
    return res.json()


@pytest.mark.anyio
async def test_create_and_get_module_definition(
    async_client: AsyncClient,
    admin_auth_headers: dict[str, str],
    leads_module: dict[str, Any],
) -> None:
    headers = admin_auth_headers
    data = leads_module
    assert data["code"] == "leads"
    assert len(data["fields"]) == 4
    assert len(data["statuses"]) == 3

    # 2. Get definition
    res_get = await async_client.get("/api/v1/modules/leads", headers=headers)
    assert res_get.status_code == 200
    get_data = res_get.json()
    assert get_data["label"] == "Leads & Prospects"
    assert get_data["record_label_singular"] == "Lead"


@pytest.mark.anyio
async def test_record_validation_and_creation(
    async_client: AsyncClient,
    admin_auth_headers: dict[str, str],
    leads_module: dict[str, Any],
) -> None:
    headers = admin_auth_headers

    # 1. Validation failure: missing required field
    res_fail = await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={"data": {"contact_name": "Acme Corp"}}, # missing title
    )
    assert res_fail.status_code == 400

    # 2. Validation failure: invalid currency/number
    res_bad_num = await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "SEO Retainer",
                "contact_name": "Acme Corp",
                "deal_value": "not-a-number",
            }
        },
    )
    assert res_bad_num.status_code == 400

    # 3. Successful creation
    res_ok = await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "SEO Retainer",
                "contact_name": "Acme Corp",
                "deal_value": 50000,
                "source": "web",
            }
        },
    )
    assert res_ok.status_code == 201
    rec = res_ok.json()
    assert rec["data"]["title"] == "SEO Retainer"
    assert rec["data"]["deal_value"] == 50000
    assert rec["status"]["code"] == "new"


@pytest.mark.anyio
async def test_record_search_and_pagination(
    async_client: AsyncClient,
    admin_auth_headers: dict[str, str],
    leads_module: dict[str, Any],
) -> None:
    headers = admin_auth_headers

    # Add leads
    await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "SEO Retainer",
                "contact_name": "Acme Corp",
                "deal_value": 50000,
                "source": "web",
            }
        },
    )
    await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "Mobile App Development",
                "contact_name": "Globex Inc",
                "deal_value": 120000,
                "source": "referral",
            }
        },
    )

    # Search for "Globex"
    res_search = await async_client.get(
        "/api/v1/modules/leads/records",
        headers=headers,
        params={"search": "Globex"},
    )
    assert res_search.status_code == 200
    search_data = res_search.json()
    assert search_data["total"] >= 1
    assert any(r["data"]["contact_name"] == "Globex Inc" for r in search_data["items"])

    # Filter by status code
    res_status = await async_client.get(
        "/api/v1/modules/leads/records",
        headers=headers,
        params={"status": "new"},
    )
    assert res_status.status_code == 200
    assert res_status.json()["total"] >= 2


@pytest.mark.anyio
async def test_record_status_transitions(
    async_client: AsyncClient,
    admin_auth_headers: dict[str, str],
    leads_module: dict[str, Any],
) -> None:
    headers = admin_auth_headers

    # Get definition to know status IDs
    res_def = await async_client.get("/api/v1/modules/leads", headers=headers)
    statuses = {s["code"]: s["id"] for s in res_def.json()["statuses"]}

    # Create record
    res_rec = await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "Branding Proposal",
                "contact_name": "Stark Industries",
                "deal_value": 75000,
            }
        },
    )
    rec_id = res_rec.json()["id"]

    # Transition to contacted
    res_trans = await async_client.post(
        f"/api/v1/modules/leads/records/{rec_id}/status",
        headers=headers,
        json={"to_status_id": statuses["contacted"], "note": "Spoke to Tony"},
    )
    assert res_trans.status_code == 200
    assert res_trans.json()["status"]["code"] == "contacted"

    # Inspect details and timeline
    res_detail = await async_client.get(
        f"/api/v1/modules/leads/records/{rec_id}",
        headers=headers,
    )
    assert res_detail.status_code == 200
    detail = res_detail.json()
    assert len(detail["activities"]) >= 2
    actions = [a["action"] for a in detail["activities"]]
    assert "created" in actions
    assert "status_changed" in actions


@pytest.mark.anyio
async def test_record_comments_and_attachments(
    async_client: AsyncClient,
    admin_auth_headers: dict[str, str],
    leads_module: dict[str, Any],
) -> None:
    headers = admin_auth_headers

    res_rec = await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "Billboard Campaign",
                "contact_name": "Wayne Enterprises",
            }
        },
    )
    rec_id = res_rec.json()["id"]

    # Add comment
    res_comment = await async_client.post(
        f"/api/v1/modules/leads/records/{rec_id}/comments",
        headers=headers,
        json={"comment": "Client requested proposal by Friday."},
    )
    assert res_comment.status_code == 201
    assert res_comment.json()["comment"] == "Client requested proposal by Friday."

    # Add attachment
    res_att = await async_client.post(
        f"/api/v1/modules/leads/records/{rec_id}/attachments",
        headers=headers,
        json={
            "file_url": "https://storage.agency.com/proposals/wayne.pdf",
            "file_name": "wayne_proposal.pdf",
            "file_size": 204800,
            "mime_type": "application/pdf",
        },
    )
    assert res_att.status_code == 201
    assert res_att.json()["file_name"] == "wayne_proposal.pdf"

    # Detail check
    res_detail = await async_client.get(
        f"/api/v1/modules/leads/records/{rec_id}",
        headers=headers,
    )
    assert res_detail.status_code == 200
    detail = res_detail.json()
    assert len(detail["comments"]) == 1
    assert len(detail["attachments"]) == 1


@pytest.mark.anyio
async def test_module_summary_aggregation(
    async_client: AsyncClient,
    admin_auth_headers: dict[str, str],
    leads_module: dict[str, Any],
) -> None:
    headers = admin_auth_headers

    # Seed some records for this test's isolated DB
    await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "SEO Retainer",
                "contact_name": "Acme Corp",
                "deal_value": 50000,
            }
        },
    )
    await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "Mobile App",
                "contact_name": "Globex",
                "deal_value": 70000,
            }
        },
    )

    res = await async_client.get("/api/v1/modules/leads/summary", headers=headers)
    assert res.status_code == 200
    data = res.json()
    assert "total_records" in data
    assert data["total_records"] >= 2
    assert "counts_by_status" in data
    assert "field_totals" in data
    assert "deal_value" in data["field_totals"]
    assert data["field_totals"]["deal_value"] == 120000


@pytest.mark.anyio
async def test_module_export_csv_and_xlsx(
    async_client: AsyncClient,
    admin_auth_headers: dict[str, str],
    leads_module: dict[str, Any],
) -> None:
    headers = admin_auth_headers

    # Seed a record for export
    await async_client.post(
        "/api/v1/modules/leads/records",
        headers=headers,
        json={
            "data": {
                "title": "Website Revamp",
                "contact_name": "Initech",
                "deal_value": 30000,
            }
        },
    )

    # 1. Export CSV
    res_csv = await async_client.get(
        "/api/v1/modules/leads/export",
        headers=headers,
        params={"format": "csv"},
    )
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]
    csv_text = res_csv.content.decode("utf-8-sig")
    assert "Lead Title" in csv_text
    assert "Contact Name" in csv_text

    # 2. Export XLSX
    res_xlsx = await async_client.get(
        "/api/v1/modules/leads/export",
        headers=headers,
        params={"format": "xlsx"},
    )
    assert res_xlsx.status_code == 200
    assert "openxmlformats" in res_xlsx.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    assert len(wb.sheetnames) >= 1
    sheet = wb.active
    assert sheet.max_row >= 2
